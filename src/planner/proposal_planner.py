"""Proposal planner — orchestrates 3-phase AI pipeline for proposal generation."""

import sys
import uuid
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from retrieval.search import search_for_proposal
from llm.gemini_client import generate_slide_plan, generate_slide_content, review_proposal
from generation.pptx_generator import generate_proposal_pptx
from generation.docx_generator import generate_proposal_docx


class ProposalSession:
    """Manages a single proposal conversation session."""

    def __init__(self):
        self.session_id = uuid.uuid4().hex[:12]
        self.history: list[dict] = []
        self.slide_plan: dict | None = None
        self.plan: dict | None = None
        self.status = "chatting"
        self.generated_files: dict = {}
        self.images: list[dict] = []
        self._references: list[dict] = []
        self._user_input: str = ""

    @property
    def reference_images(self) -> list[dict]:
        return [img for img in self.images if img.get("type") == "reference"]

    @property
    def embed_images(self) -> list[dict]:
        return [img for img in self.images if img.get("type") == "embed"]

    def chat(self, user_message: str) -> dict:
        self.history.append({"role": "user", "content": user_message})
        self._user_input = user_message

        references = self._find_references(user_message)
        if references:
            self._references = references.get("content_references", [])

        try:
            result = generate_slide_plan(
                user_input=user_message,
                references=self._references if self._references else [],
                conversation_history=self.history,
                images=self.reference_images if self.reference_images else None,
            )
        except Exception as e:
            error_msg = f"I encountered an issue processing your request: {str(e)}"
            self.history.append({"role": "assistant", "content": error_msg})
            return {"ready": False, "message": error_msg}

        if result.get("ready"):
            self.slide_plan = result
            self.status = "plan_ready"

            slide_count = sum(
                len(s.get("slides", [1])) for s in result.get("slide_plan", [])
            )

            self.history.append({
                "role": "assistant",
                "content": f"I've prepared a proposal plan: **{result.get('title', 'Proposal')}** for **{result.get('customer', 'your client')}** with {slide_count} planned slides. Ready to generate the documents!",
            })
            return {
                "ready": True,
                "message": self.history[-1]["content"],
                "plan_summary": {
                    "title": result.get("title", ""),
                    "customer": result.get("customer", ""),
                    "objective": result.get("objective", ""),
                    "sections": len(result.get("slide_plan", [])),
                },
            }
        else:
            msg = result.get("message", "Could you provide more details?")
            self.history.append({"role": "assistant", "content": msg})
            return {"ready": False, "message": msg}

    def generate(self, template_name: str | None = None) -> dict:
        if not self.slide_plan:
            return {"error": "No proposal plan ready. Keep chatting to build one."}

        self.status = "generating"
        output_dir = Path(__file__).resolve().parent.parent.parent / "outputs"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = self.slide_plan.get("customer", "proposal").replace(" ", "_")[:30]

        results = {}

        try:
            content_plan = generate_slide_content(
                slide_plan=self.slide_plan,
                user_input=self._user_input,
                references=self._references if self._references else None,
            )

            self.plan = content_plan

            pptx_path = output_dir / "ppt" / f"{safe_title}_{timestamp}.pptx"
            generate_proposal_pptx(content_plan, pptx_path, template_name=template_name,
                                   embed_images=self.embed_images if self.embed_images else None)

            review = self._review_and_patch(content_plan, pptx_path, template_name)
            if review and review.get("patched"):
                results["review_score"] = review.get("score", "N/A")
                results["review_improved"] = True
            elif review:
                results["review_score"] = review.get("score", "N/A")

            results["pptx"] = str(pptx_path)
            self.generated_files["pptx"] = str(pptx_path)
        except Exception as e:
            results["pptx_error"] = str(e)

        try:
            docx_path = output_dir / "docx" / f"{safe_title}_{timestamp}.docx"
            plan_for_docx = self.plan or self.slide_plan
            generate_proposal_docx(plan_for_docx, docx_path)
            results["docx"] = str(docx_path)
            self.generated_files["docx"] = str(docx_path)
        except Exception as e:
            results["docx_error"] = str(e)

        self.status = "complete"
        return results

    def _review_and_patch(self, content_plan: dict, pptx_path, template_name: str | None) -> dict | None:
        try:
            from pptx import Presentation
            prs = Presentation(str(pptx_path))
            slide_texts = []
            for slide in prs.slides:
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                texts.append(text)
                slide_texts.append("\n".join(texts) if texts else "(empty slide)")

            review = review_proposal(
                user_input=self._user_input,
                plan=content_plan,
                slide_texts=slide_texts,
            )

            score = review.get("score", 10)
            if score < 7 and review.get("content_patches"):
                sections = content_plan.get("sections", {})
                for section_key, patch in review["content_patches"].items():
                    if section_key in sections:
                        if isinstance(sections[section_key], dict) and isinstance(patch, dict):
                            sections[section_key].update(patch)

                generate_proposal_pptx(content_plan, pptx_path, template_name=template_name,
                                       embed_images=self.embed_images if self.embed_images else None)
                self.plan = content_plan
                return {"score": score, "patched": True}

            return {"score": score, "patched": False}
        except Exception:
            return None

    def _find_references(self, query: str) -> dict | None:
        try:
            return search_for_proposal(query, top_k=10)
        except Exception:
            return None
