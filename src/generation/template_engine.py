"""Template Engine — loads Xebia base template and provides layout utilities.

The template provides slide masters with branded elements, placeholder
positioning, and canvas geometry. Visual identity (colors, fonts, accents)
is generated dynamically per proposal by design_generator.py.

Canvas: 13.333 x 7.500 inches (widescreen 16:9)
"""

from __future__ import annotations

import random
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "templates"
FROM_DOCS_DIR = TEMPLATES_DIR / "from_documents"
TEMPLATE_FILE = "xebia_retail.pptx"

TEMPLATE_REGISTRY: dict[str, dict] = {
    "xebia_retail": {
        "file": "xebia_retail.pptx",
        "dir": TEMPLATES_DIR,
        "family": "xebia_standard",
        "display_name": "Xebia Standard",
        "description": "Clean, modern Xebia branded template — best for most proposals",
        "content_layout": "Content_Basic",
        "cover_layout": "Main-cover_dark",
        "chapter_layout": "Chapter_light",
        "closing_layout": "End-cover_dark",
        "toc_layout": "Table of contents",
    },
    "xebia_synapse": {
        "file": "Proposal-Synapse_to_Fabric_Migration.pptx",
        "dir": FROM_DOCS_DIR,
        "family": "xebia_standard",
        "display_name": "Synapse Migration",
        "description": "Data platform migration theme — ideal for Fabric/Synapse projects",
        "content_layout": "Content_Basic",
        "cover_layout": "Main-cover_dark",
        "chapter_layout": "Chapter_light",
        "closing_layout": "End-cover_dark",
        "toc_layout": "Table of contents",
    },
    "xebia_pnb": {
        "file": "Xebias_Technical_Proposal_for_PNB_MetLife.pptx",
        "dir": FROM_DOCS_DIR,
        "family": "xebia_standard",
        "display_name": "Enterprise Technical",
        "description": "Technical depth template — suited for insurance, BFSI proposals",
        "content_layout": "Content_Basic",
        "cover_layout": "Main-cover_dark",
        "chapter_layout": "Chapter_light",
        "closing_layout": "End-cover_dark",
        "toc_layout": "Table of contents",
    },
    "xebia_mohesr": {
        "file": "Xebias_Microsoft_Fabric_Data_Platform_Implementation_for_MOH.pptx",
        "dir": FROM_DOCS_DIR,
        "family": "xebia_standard",
        "display_name": "Fabric Platform",
        "description": "Microsoft Fabric data platform focused — government & enterprise",
        "content_layout": "Content_Basic",
        "cover_layout": "Main-cover_dark",
        "chapter_layout": "Chapter_light",
        "closing_layout": "End-cover_dark",
        "toc_layout": "Table of contents",
    },
    "xebia_carrington": {
        "file": "Xebia-Carrington_Microsoft_Fabric_Implementation_Engagement_.pptx",
        "dir": FROM_DOCS_DIR,
        "family": "carrington",
        "display_name": "Carrington Style",
        "description": "Co-branded style template — for joint-venture or partner proposals",
        "content_layout": "Title Content/White",
        "cover_layout": "Title Slide / Dark",
        "chapter_layout": "Section Divider / Light",
        "closing_layout": "Title Slide / Dark",
        "toc_layout": "Index Slides",
    },
    "xebia_hct": {
        "file": "Xebias_Proposal_Walkthrough_for_HCT.pptx",
        "dir": FROM_DOCS_DIR,
        "family": "hct",
        "display_name": "Executive Walkthrough",
        "description": "Executive-level walkthrough — minimal text, high visual impact",
        "content_layout": "Content_01_Left",
        "cover_layout": "Cover_01_Dark",
        "chapter_layout": "Section_Slides_01",
        "closing_layout": "Cover_01_Dark",
        "toc_layout": "Intro_01",
    },
    "xebia_retail_cases": {
        "file": "Data_&_AI__Retail_Case_Studies_Use_Case_Aligned.pptx",
        "dir": FROM_DOCS_DIR,
        "family": "xebia_standard",
        "display_name": "Retail & Case Studies",
        "description": "Case-study heavy template — great for retail, CPG, FMCG proposals",
        "content_layout": "Content_Basic",
        "cover_layout": "Main-cover_dark",
        "chapter_layout": "Chapter_light",
        "closing_layout": "End-cover_dark",
        "toc_layout": "Table of contents",
    },
}


def get_template_info() -> list[dict]:
    """Return metadata for all available templates."""
    result = []
    for name, info in TEMPLATE_REGISTRY.items():
        path = info["dir"] / info["file"]
        if path.exists():
            result.append({
                "name": name,
                "display_name": info.get("display_name", name),
                "family": info.get("family", "xebia_standard"),
                "description": info.get("description", ""),
            })
    return result


def get_available_templates() -> list[str]:
    """Return names of templates whose files actually exist."""
    available = []
    for name, info in TEMPLATE_REGISTRY.items():
        path = info["dir"] / info["file"]
        if path.exists():
            available.append(name)
    return available


def pick_template(template_name: str | None = None) -> dict:
    """Pick a template by name, or randomly from available ones."""
    if template_name and template_name in TEMPLATE_REGISTRY:
        info = TEMPLATE_REGISTRY[template_name]
        path = info["dir"] / info["file"]
        if path.exists():
            return {**info, "name": template_name, "path": str(path)}

    available = get_available_templates()
    if not available:
        return {
            "name": "xebia_retail",
            "path": str(TEMPLATES_DIR / TEMPLATE_FILE),
            "family": "xebia_standard",
            "content_layout": "Content_Basic",
            "cover_layout": "Main-cover_dark",
            "chapter_layout": "Chapter_light",
            "closing_layout": "End-cover_dark",
            "toc_layout": "Table of contents",
        }

    chosen = random.choice(available)
    info = TEMPLATE_REGISTRY[chosen]
    return {**info, "name": chosen, "path": str(info["dir"] / info["file"])}


def generate_theme():
    """Return template path for loading."""
    return str(TEMPLATES_DIR / TEMPLATE_FILE)


def create_themed_presentation(template_path=None) -> Presentation:
    if template_path is None:
        template_path = generate_theme()
    tpl = Path(template_path) if isinstance(template_path, str) else template_path
    if tpl.exists() if isinstance(tpl, Path) else Path(tpl).exists():
        return Presentation(str(tpl))
    return Presentation()


def _rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_world_map(slide, color: str = "#6C1D5F"):
    """Add world map dot visualization for Xebia's global presence."""
    c = _rgb(color)

    world_dots = [
        (4.2, 2.8), (4.5, 2.6), (4.8, 2.4), (5.0, 2.5), (5.3, 2.3),
        (5.5, 2.5), (5.8, 2.4), (6.0, 2.6), (6.3, 2.5), (6.5, 2.3),
        (6.8, 2.5), (7.0, 2.7), (7.3, 2.6), (7.5, 2.8), (7.8, 2.7),
        (8.0, 2.9), (8.3, 3.0), (8.5, 2.8), (8.8, 3.1), (9.0, 3.0),
        (4.0, 3.0), (4.3, 3.2), (4.5, 3.4), (4.8, 3.3), (5.0, 3.5),
        (5.3, 3.4), (5.5, 3.6), (5.8, 3.5), (6.0, 3.7), (6.3, 3.6),
        (6.5, 3.8), (6.8, 3.7), (7.0, 3.9), (7.3, 3.8), (7.5, 4.0),
        (3.5, 3.5), (3.8, 3.7), (4.0, 3.9), (4.3, 4.0), (4.5, 4.2),
        (9.2, 3.2), (9.4, 3.4), (9.6, 3.3), (9.8, 3.5), (9.5, 3.6),
        (9.0, 4.0), (9.3, 4.2), (9.5, 4.4), (9.7, 4.3), (9.9, 4.5),
        (7.8, 4.2), (8.0, 4.4), (8.2, 4.3), (8.5, 4.5), (8.7, 4.4),
        (3.0, 4.0), (3.2, 4.2), (3.5, 4.1), (3.7, 4.3), (3.3, 4.5),
        (5.0, 2.0), (5.5, 1.8), (6.0, 2.0), (6.5, 1.9), (7.0, 2.1),
        (7.5, 2.0), (8.0, 2.2), (8.5, 2.1),
    ]

    xebia_offices = [
        (5.2, 2.8), (5.4, 2.6), (5.6, 3.0), (7.5, 3.2),
        (8.0, 3.5), (3.5, 3.0), (9.0, 4.8), (6.2, 3.8),
    ]

    r, g, b = c[0], c[1], c[2]
    faded = RGBColor(
        min(255, int(r + (255 - r) * 0.7)),
        min(255, int(g + (255 - g) * 0.7)),
        min(255, int(b + (255 - b) * 0.7)),
    )

    for x, y in world_dots:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                     Inches(0.07), Inches(0.07))
        dot.fill.solid()
        dot.fill.fore_color.rgb = faded
        dot.line.fill.background()

    pulse_color = RGBColor(
        min(255, int(r + (255 - r) * 0.6)),
        min(255, int(g + (255 - g) * 0.6)),
        min(255, int(b + (255 - b) * 0.6)),
    )

    for x, y in xebia_offices:
        pulse = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.04), Inches(y - 0.04),
                                       Inches(0.22), Inches(0.22))
        pulse.fill.solid()
        pulse.fill.fore_color.rgb = pulse_color
        pulse.line.fill.background()

        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                     Inches(0.14), Inches(0.14))
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.fill.background()
