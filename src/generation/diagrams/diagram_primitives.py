"""Low-level diagram drawing primitives — arrows, dashed containers,
icon placeholders, labels, and flow connectors."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from generation.design_engine.primitives import (
    rgb, set_dash_style, add_rectangle, add_rounded_rectangle, add_circle,
)


SLIDE_W = 13.333
SLIDE_H = 7.500


def add_textbox(slide, left, top, width, height, text, size, color,
                bold=False, alignment=None, font_name="Arial"):
    left = max(0, min(left, SLIDE_W - 0.1))
    top = max(0, min(top, SLIDE_H - 0.1))
    width = min(width, SLIDE_W - left)
    height = min(height, SLIDE_H - top)
    tx = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(text)
    p.font.size = Pt(size)
    p.font.color.rgb = rgb(color) if isinstance(color, str) else color
    p.font.bold = bold
    p.font.name = font_name
    if alignment:
        p.alignment = alignment
    return tx


def add_dashed_container(slide, left, top, width, height,
                         label="", label_color="#333333",
                         border_color="#6C1D5F", dash="dash",
                         fill_color=None, border_width=1.5,
                         font_name="Arial"):
    """Draw a dashed-border container with an optional top-left label."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill_color)
    else:
        shape.fill.background()
    set_dash_style(shape, dash=dash, color=border_color, width=border_width)

    if label:
        add_textbox(slide, left + 0.1, top + 0.05, width - 0.2, 0.25,
                    label, 8, label_color, bold=True, font_name=font_name)
    return shape


def add_icon_placeholder(slide, x, y, size, label_char, bg_color,
                         text_color="#FFFFFF"):
    """Colored circle with 1-2 char abbreviation as icon stand-in."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x), Inches(y), Inches(size), Inches(size),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(bg_color)
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = label_char[:2]
    p.font.size = Pt(int(size * 26))
    p.font.color.rgb = rgb(text_color)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_horizontal_arrow(slide, x1, y, x2, color="#C00000", thickness=0.03):
    """Draw a horizontal arrow from x1 to x2 at vertical position y."""
    if abs(x2 - x1) < 0.2:
        return

    going_right = x2 > x1
    line_start = min(x1, x2)
    line_end = max(x1, x2) - 0.15
    line_w = line_end - line_start

    if line_w > 0.05:
        add_rectangle(slide, line_start, y - thickness / 2,
                      line_w, thickness, fill_color=color)

    head_x = (max(x1, x2) - 0.15) if going_right else min(x1, x2)
    head = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(head_x), Inches(y - 0.06), Inches(0.15), Inches(0.12),
    )
    head.fill.solid()
    head.fill.fore_color.rgb = rgb(color)
    head.line.fill.background()
    head.rotation = 90 if going_right else 270
    return head


def add_vertical_arrow(slide, x, y1, y2, color="#C00000", thickness=0.03):
    """Draw a vertical arrow from y1 to y2 at horizontal position x."""
    if abs(y2 - y1) < 0.2:
        return

    going_down = y2 > y1
    line_start = min(y1, y2)
    line_end = max(y1, y2) - 0.15
    line_h = line_end - line_start

    if line_h > 0.05:
        add_rectangle(slide, x - thickness / 2, line_start,
                      thickness, line_h, fill_color=color)

    head_y = (max(y1, y2) - 0.15) if going_down else min(y1, y2)
    head = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(x - 0.06), Inches(head_y), Inches(0.12), Inches(0.15),
    )
    head.fill.solid()
    head.fill.fore_color.rgb = rgb(color)
    head.line.fill.background()
    head.rotation = 180 if going_down else 0
    return head


def add_labeled_box(slide, x, y, w, h, label, sublabel="",
                    fill_color="#FFFFFF", border_color="#6C1D5F",
                    text_color="#333333", border_width=1.0,
                    font_name="Arial"):
    """A rounded rect with centered label and optional sublabel."""
    shape = add_rounded_rectangle(slide, x, y, w, h,
                                  fill_color=fill_color,
                                  line_color=border_color,
                                  line_width=border_width)

    label_h = 0.22 if sublabel else h - 0.08
    label_y = y + 0.04 if sublabel else y + (h - label_h) / 2
    add_textbox(slide, x + 0.04, label_y, w - 0.08, label_h,
                label, 8, text_color, bold=True,
                alignment=PP_ALIGN.CENTER, font_name=font_name)

    if sublabel:
        sub_y = label_y + 0.20
        sub_h = h - 0.30
        if sub_h > 0.1:
            add_textbox(slide, x + 0.04, sub_y, w - 0.08, sub_h,
                        sublabel, 7, "#666666",
                        alignment=PP_ALIGN.CENTER, font_name=font_name)
    return shape


def add_band(slide, left, top, width, height, label, items=None,
             description="", fill_color="#F5EFF5", text_color="#333333",
             accent_color="#6C1D5F", font_name="Arial"):
    """Horizontal band with label on left and items/description on right."""
    add_rounded_rectangle(slide, left, top, width, height,
                          fill_color=fill_color)

    label_w = min(1.8, width * 0.2)
    add_rounded_rectangle(slide, left + 0.04, top + 0.04,
                          label_w, height - 0.08,
                          fill_color=accent_color)
    add_textbox(slide, left + 0.08, top + (height - 0.20) / 2,
                label_w - 0.08, 0.20, label, 7, "#FFFFFF",
                bold=True, font_name=font_name)

    items_x = left + label_w + 0.15
    items_w = width - label_w - 0.25
    if items:
        items_text = "  ·  ".join(items)
        add_textbox(slide, items_x, top + 0.04, items_w, height * 0.45,
                    items_text, 7, text_color, font_name=font_name)
    if description:
        desc_y = top + (height * 0.45 if items else 0.04)
        add_textbox(slide, items_x, desc_y, items_w, height * 0.45,
                    description, 6, "#666666", font_name=font_name)
