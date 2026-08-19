"""Xebia Proposal DOCX Generator.

Takes a structured proposal plan (dict/YAML) and generates a complete
Xebia-branded Word document with professional formatting, headers,
footers, page numbers, and table of contents.
"""

from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from design_system.brand import colors, typography, spacing, theme, copyright_text


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _brand_purple() -> RGBColor:
    return _rgb(colors()["primary"]["purple"])


class ProposalDOCXGenerator:
    """Generates a complete Xebia-branded DOCX from a proposal plan."""

    def __init__(self, plan: dict):
        self.plan = plan
        self.doc = Document()
        self._setup_styles()
        self._setup_page()

    def generate(self, output_path: str | Path) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        storyline = self.plan.get("storyline", [])
        sections = self.plan.get("sections", {})

        for section_key in storyline:
            section_data = sections.get(section_key, {})
            builder = getattr(self, f"_build_{section_key}", self._build_generic_section)
            builder(section_key, section_data)

        self._add_headers_footers()
        self.doc.save(str(output_path))
        return output_path

    def _setup_styles(self):
        """Configure document styles to match Xebia brand."""
        typo = typography()["docx"]
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = typo["body"]["font"]
        font.size = Pt(typo["body"]["size_pt"])
        font.color.rgb = _rgb(colors()["docx"]["body_text"])

        pf = style.paragraph_format
        pf.space_after = Pt(spacing()["docx"]["paragraph_spacing_pt"])
        pf.line_spacing = typo["body"]["line_spacing"]

        for level in range(1, 4):
            heading_key = f"heading{level}"
            if heading_key not in typo:
                continue
            h_config = typo[heading_key]
            h_style = self.doc.styles[f"Heading {level}"]
            h_style.font.name = h_config["font"]
            h_style.font.size = Pt(h_config["size_pt"])
            h_style.font.bold = h_config["weight"] == "bold"
            h_style.font.color.rgb = _rgb(colors()["docx"].get(
                "accent_text" if level == 3 else "heading_color",
                colors()["neutral"]["black"]
            ))
            h_style.paragraph_format.space_before = Pt(
                spacing()["docx"]["heading_spacing_before_pt"]
            )
            h_style.paragraph_format.space_after = Pt(
                spacing()["docx"]["heading_spacing_after_pt"]
            )

    def _setup_page(self):
        """Set page margins."""
        sp = spacing()["docx"]
        section = self.doc.sections[0]
        section.top_margin = Inches(sp["margin_top_in"])
        section.bottom_margin = Inches(sp["margin_bottom_in"])
        section.left_margin = Inches(sp["margin_left_in"])
        section.right_margin = Inches(sp["margin_right_in"])

    def _add_headers_footers(self):
        """Add headers and footers to all sections."""
        for section in self.doc.sections:
            section.different_first_page_header_footer = True

            # Regular header — thin purple line
            header = section.header
            header.is_linked_to_previous = False
            hp = header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = hp.add_run("Xebia  |  ")
            run.font.size = Pt(9)
            run.font.color.rgb = _brand_purple()
            run.font.bold = True
            run.font.name = "Arial"
            run2 = hp.add_run(self.plan.get("title", "Proposal"))
            run2.font.size = Pt(9)
            run2.font.color.rgb = _rgb(colors()["neutral"]["mid_gray"])
            run2.font.name = "Calibri"

            # Add bottom border to header
            pPr = hp._element.get_or_add_pPr()
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'  <w:bottom w:val="single" w:sz="4" w:space="4" w:color="{colors()["primary"]["purple"].lstrip("#")}"/>'
                f'</w:pBdr>'
            )
            pPr.append(pBdr)

            # Footer — copyright left, page number right
            footer = section.footer
            footer.is_linked_to_previous = False
            fp = footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run_copy = fp.add_run(copyright_text())
            run_copy.font.size = Pt(8)
            run_copy.font.color.rgb = _rgb(colors()["neutral"]["mid_gray"])
            run_copy.font.name = "Calibri"

            # Tab to right + page number
            run_tab = fp.add_run("\t\t")
            run_tab.font.size = Pt(8)

            run_page = fp.add_run("Page ")
            run_page.font.size = Pt(8)
            run_page.font.color.rgb = _rgb(colors()["neutral"]["mid_gray"])
            run_page.font.name = "Calibri"

            # Auto page number field
            fld_xml = parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
            )
            fp._element.append(fld_xml)

    # ============================================================
    # SECTION BUILDERS
    # ============================================================

    def _build_cover(self, key: str, data: dict):
        """Cover page with centered title and purple accent."""
        # Add some spacing at top
        for _ in range(6):
            self.doc.add_paragraph("")

        # Title
        title = data.get("title", self.plan.get("title", "Proposal"))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.size = Pt(32)
        run.font.color.rgb = _brand_purple()
        run.font.bold = True
        run.font.name = "Arial"

        # Purple divider line
        p_line = self.doc.add_paragraph()
        p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_line = p_line.add_run("━" * 30)
        run_line.font.color.rgb = _brand_purple()
        run_line.font.size = Pt(12)

        # Subtitle
        subtitle = data.get("subtitle", self.plan.get("objective", ""))
        if subtitle:
            p_sub = self.doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_sub = p_sub.add_run(subtitle)
            run_sub.font.size = Pt(14)
            run_sub.font.color.rgb = _rgb(colors()["neutral"]["dark_gray"])
            run_sub.font.name = "Calibri"

        # Customer and date
        self.doc.add_paragraph("")
        customer = data.get("customer", self.plan.get("customer", ""))
        if customer:
            p_cust = self.doc.add_paragraph()
            p_cust.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cust = p_cust.add_run(f"Prepared for: {customer}")
            run_cust.font.size = Pt(12)
            run_cust.font.color.rgb = _rgb(colors()["neutral"]["mid_gray"])
            run_cust.font.name = "Calibri"

        p_date = self.doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_date = p_date.add_run(data.get("date", datetime.now().strftime("%B %Y")))
        run_date.font.size = Pt(11)
        run_date.font.color.rgb = _rgb(colors()["neutral"]["mid_gray"])
        run_date.font.name = "Calibri"

        # Tagline at bottom
        for _ in range(4):
            self.doc.add_paragraph("")
        p_tag = self.doc.add_paragraph()
        p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_tag = p_tag.add_run(theme()["tagline"])
        run_tag.font.size = Pt(11)
        run_tag.font.color.rgb = _brand_purple()
        run_tag.font.italic = True
        run_tag.font.name = "Calibri"

        self.doc.add_page_break()

    def _build_table_of_contents(self, key: str, data: dict):
        """Add a Table of Contents placeholder."""
        self.doc.add_heading("Table of Contents", level=1)

        # TOC field — Word will populate this when user presses Update
        p = self.doc.add_paragraph()
        run = p.add_run()
        fld_begin = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>')
        fld_instr = parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>')
        fld_sep = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>')
        fld_text = parse_xml(f'<w:r {nsdecls("w")}><w:t>[Right-click and select Update Field to generate TOC]</w:t></w:r>')
        fld_end = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')

        p._element.append(fld_begin)
        p._element.append(fld_instr)
        p._element.append(fld_sep)
        p._element.append(fld_text)
        p._element.append(fld_end)

        self.doc.add_page_break()

    def _build_executive_summary(self, key: str, data: dict):
        """Executive summary section."""
        self.doc.add_heading(data.get("title", "Executive Summary"), level=1)

        if data.get("summary"):
            p = self.doc.add_paragraph(data["summary"])
            p.paragraph_format.space_after = Pt(12)

        if data.get("key_points"):
            self.doc.add_heading("Key Highlights", level=2)
            for point in data["key_points"]:
                p = self.doc.add_paragraph(point, style="List Bullet")

    def _build_proposed_solution(self, key: str, data: dict):
        """Solution section with optional sub-sections."""
        title = data.get("title", "Proposed Solution")
        self.doc.add_heading(title, level=1)

        if data.get("body"):
            self.doc.add_paragraph(data["body"])

        if data.get("left") and data.get("right"):
            self.doc.add_heading(data["left"].get("title", ""), level=2)
            for bullet in data["left"].get("bullets", []):
                self.doc.add_paragraph(bullet, style="List Bullet")

            self.doc.add_heading(data["right"].get("title", ""), level=2)
            for bullet in data["right"].get("bullets", []):
                self.doc.add_paragraph(bullet, style="List Bullet")

        if data.get("bullets"):
            for bullet in data["bullets"]:
                self.doc.add_paragraph(bullet, style="List Bullet")

    def _build_timeline(self, key: str, data: dict):
        """Timeline section with a table."""
        self.doc.add_heading(data.get("title", "Timeline"), level=1)

        phases = data.get("phases", [])
        if not phases:
            return

        table = self.doc.add_table(rows=len(phases) + 1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        headers = ["Phase", "Duration", "Description"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            self._style_table_header_cell(cell)

        for r_idx, phase in enumerate(phases):
            table.rows[r_idx + 1].cells[0].text = phase.get("name", "")
            table.rows[r_idx + 1].cells[1].text = phase.get("duration", "")
            table.rows[r_idx + 1].cells[2].text = phase.get("description", "")

            if r_idx % 2 == 1:
                for cell in table.rows[r_idx + 1].cells:
                    self._shade_cell(cell, colors()["neutral"]["off_white"])

    def _build_team_structure(self, key: str, data: dict):
        """Team structure as a table."""
        self.doc.add_heading(data.get("title", "Team Structure"), level=1)

        members = data.get("members", [])
        if not members:
            return

        table = self.doc.add_table(rows=len(members) + 1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        headers = ["Name", "Role", "Expertise"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            self._style_table_header_cell(cell)

        for r_idx, member in enumerate(members):
            table.rows[r_idx + 1].cells[0].text = member.get("name", "")
            table.rows[r_idx + 1].cells[1].text = member.get("role", "")
            table.rows[r_idx + 1].cells[2].text = member.get("expertise", "")

            if r_idx % 2 == 1:
                for cell in table.rows[r_idx + 1].cells:
                    self._shade_cell(cell, colors()["neutral"]["off_white"])

    def _build_commercials(self, key: str, data: dict):
        """Commercials section with pricing table."""
        self.doc.add_heading(data.get("title", "Commercials"), level=1)

        rows = data.get("rows", [])
        if not rows:
            return

        table = self.doc.add_table(rows=len(rows) + 1 + (1 if data.get("total") else 0), cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        headers = ["Item", "Hours", "Rate", "Cost"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            self._style_table_header_cell(cell)

        for r_idx, row in enumerate(rows):
            table.rows[r_idx + 1].cells[0].text = row.get("item", "")
            table.rows[r_idx + 1].cells[1].text = str(row.get("hours", ""))
            table.rows[r_idx + 1].cells[2].text = str(row.get("rate", ""))
            table.rows[r_idx + 1].cells[3].text = str(row.get("cost", ""))

        if data.get("total"):
            total_row_idx = len(rows) + 1
            table.rows[total_row_idx].cells[0].text = "Total"
            table.rows[total_row_idx].cells[3].text = data["total"]
            for cell in table.rows[total_row_idx].cells:
                self._style_table_header_cell(cell)

        if data.get("assumptions"):
            self.doc.add_paragraph("")
            self.doc.add_heading("Assumptions", level=2)
            for assumption in data["assumptions"]:
                self.doc.add_paragraph(assumption, style="List Bullet")

    def _build_closing(self, key: str, data: dict):
        """Closing section with next steps and contact."""
        self.doc.add_heading(data.get("title", "Next Steps"), level=1)

        if data.get("body"):
            self.doc.add_paragraph(data["body"])

        if data.get("bullets"):
            for bullet in data["bullets"]:
                self.doc.add_paragraph(bullet, style="List Bullet")

        if data.get("contact_name"):
            self.doc.add_paragraph("")
            p = self.doc.add_paragraph()
            run = p.add_run("Contact: ")
            run.font.bold = True
            run.font.color.rgb = _brand_purple()
            run2 = p.add_run(data["contact_name"])
            if data.get("contact_email"):
                run3 = p.add_run(f"  |  {data['contact_email']}")
                run3.font.color.rgb = _rgb(colors()["neutral"]["mid_gray"])

    def _build_generic_section(self, key: str, data: dict):
        """Fallback builder that extracts content from any layout type."""
        title = data.get("title", key.replace("_", " ").title())
        self.doc.add_heading(title, level=1)

        if data.get("summary"):
            self.doc.add_paragraph(data["summary"])
        if data.get("body"):
            self.doc.add_paragraph(data["body"])
        if data.get("bullets"):
            for bullet in data["bullets"]:
                self.doc.add_paragraph(bullet, style="List Bullet")

        self._render_layout_content(data)

        if data.get("slides"):
            for slide_data in data["slides"]:
                if slide_data.get("title"):
                    self.doc.add_heading(slide_data["title"], level=2)
                if slide_data.get("body"):
                    self.doc.add_paragraph(slide_data["body"])
                if slide_data.get("bullets"):
                    for bullet in slide_data["bullets"]:
                        self.doc.add_paragraph(bullet, style="List Bullet")
                self._render_layout_content(slide_data)

    def _render_layout_content(self, data: dict):
        """Extract and render content from any slide layout type into DOCX."""
        layout = data.get("layout", "")

        if data.get("items") and not data.get("bullets"):
            for item in data["items"]:
                p = self.doc.add_paragraph(style="List Bullet")
                run = p.add_run(f"{item.get('label', '')}: ")
                run.font.bold = True
                p.add_run(item.get("description", ""))

        if data.get("steps"):
            for i, step in enumerate(data["steps"], 1):
                p = self.doc.add_paragraph()
                run = p.add_run(f"{i}. {step.get('label', '')}: ")
                run.font.bold = True
                p.add_run(step.get("description", ""))

        if data.get("stats"):
            for stat in data["stats"]:
                p = self.doc.add_paragraph(style="List Bullet")
                run = p.add_run(f"{stat.get('value', '')} ")
                run.font.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = _brand_purple()
                p.add_run(f"— {stat.get('label', '')}")

        if data.get("pairs"):
            table = self.doc.add_table(rows=len(data["pairs"]), cols=2)
            table.style = "Table Grid"
            for i, pair in enumerate(data["pairs"]):
                table.rows[i].cells[0].text = pair.get("key", "")
                table.rows[i].cells[1].text = pair.get("value", "")
                for p in table.rows[i].cells[0].paragraphs:
                    for run in p.runs:
                        run.font.bold = True

        if layout == "two_column":
            left = data.get("left", {})
            right = data.get("right", {})
            if left:
                self.doc.add_heading(left.get("title", ""), level=3)
                for b in left.get("bullets", []):
                    self.doc.add_paragraph(b, style="List Bullet")
            if right:
                self.doc.add_heading(right.get("title", ""), level=3)
                for b in right.get("bullets", []):
                    self.doc.add_paragraph(b, style="List Bullet")

        if data.get("layers"):
            self.doc.add_heading("Architecture Layers", level=3)
            for layer in data["layers"]:
                p = self.doc.add_paragraph(style="List Bullet")
                run = p.add_run(f"{layer.get('name', '')}: ")
                run.font.bold = True
                p.add_run(", ".join(layer.get("components", [])))

        if data.get("technologies"):
            table = self.doc.add_table(rows=len(data["technologies"]) + 1, cols=3)
            table.style = "Table Grid"
            for i, h in enumerate(["Technology", "Category", "Description"]):
                cell = table.rows[0].cells[i]
                cell.text = h
                self._style_table_header_cell(cell)
            for i, tech in enumerate(data["technologies"]):
                table.rows[i + 1].cells[0].text = tech.get("name", "")
                table.rows[i + 1].cells[1].text = tech.get("category", "")
                table.rows[i + 1].cells[2].text = tech.get("description", "")

        if data.get("challenges"):
            table = self.doc.add_table(rows=len(data["challenges"]) + 1, cols=3)
            table.style = "Table Grid"
            for i, h in enumerate(["Challenge", "Impact", "Mitigation"]):
                cell = table.rows[0].cells[i]
                cell.text = h
                self._style_table_header_cell(cell)
            for i, ch in enumerate(data["challenges"]):
                table.rows[i + 1].cells[0].text = ch.get("challenge", "")
                table.rows[i + 1].cells[1].text = ch.get("impact", "")
                table.rows[i + 1].cells[2].text = ch.get("solution", "")

        if layout == "comparison_table":
            headers = data.get("headers", [])
            rows = data.get("rows", [])
            if headers and rows:
                table = self.doc.add_table(rows=len(rows) + 1, cols=len(headers))
                table.style = "Table Grid"
                for i, h in enumerate(headers):
                    cell = table.rows[0].cells[i]
                    cell.text = h
                    self._style_table_header_cell(cell)
                for r_idx, row in enumerate(rows):
                    for c_idx, val in enumerate(row):
                        if c_idx < len(headers):
                            table.rows[r_idx + 1].cells[c_idx].text = str(val)

    # ============================================================
    # TABLE STYLING HELPERS
    # ============================================================

    def _style_table_header_cell(self, cell):
        """Apply purple header styling to a table cell."""
        self._shade_cell(cell, colors()["primary"]["purple"])
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = _rgb("#FFFFFF")
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.name = "Arial"

    def _shade_cell(self, cell, hex_color: str):
        """Apply background shading to a cell."""
        hex_clean = hex_color.lstrip("#")
        shading = parse_xml(
            f'<w:shd {nsdecls("w")} w:fill="{hex_clean}" w:val="clear"/>'
        )
        cell._element.get_or_add_tcPr().append(shading)


def generate_proposal_docx(plan: dict, output_path: str | Path) -> Path:
    """Convenience function to generate a DOCX from a plan dict."""
    generator = ProposalDOCXGenerator(plan)
    return generator.generate(output_path)


if __name__ == "__main__":
    import yaml

    if len(sys.argv) < 3:
        print("Usage: python docx_generator.py <plan.yaml> <output.docx>")
        sys.exit(1)

    with open(sys.argv[1], encoding="utf-8") as f:
        plan = yaml.safe_load(f)

    out = generate_proposal_docx(plan, sys.argv[2])
    print(f"Generated: {out}")
